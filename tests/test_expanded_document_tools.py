"""
tests/test_expanded_document_tools.py
Comprehensive automated test suite for the expanded OFFAIR Document Tools Suite.
Zero AI, zero Ollama, zero Docker, purely local and air-gap verified.
"""

import io
import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from document_tools.pdf_ops import (
    add_blank_page,
    add_page_numbers,
    add_visual_signature,
    add_watermark,
    compress_pdf,
    crop_pdf,
    decrypt_pdf,
    duplicate_page,
    encrypt_pdf,
    extract_pages,
    extract_text_from_pdf,
    get_pdf_metadata,
    images_to_pdf,
    inspect_pdf_security,
    merge_pdfs,
    pdf_to_docx,
    pdf_to_markdown,
    remove_pdf_metadata,
    render_pdf_to_images,
    reorder_pages,
    resize_pdf,
    rotate_pdf,
    split_pdf,
)
from document_tools.office_ops import (
    csv_to_xlsx,
    docx_extract_tables,
    docx_get_stats,
    docx_to_markdown,
    docx_to_pdf,
    docx_to_text,
    pptx_extract_text,
    pptx_get_stats,
    spreadsheet_preview,
    txt_to_docx,
    txt_to_pdf,
    xlsx_to_csv,
)
from document_tools.image_ops import (
    batch_convert_images,
    compress_image,
    convert_image,
    crop_image,
    resize_image,
    rotate_flip_image,
)
from document_tools.inspector_ops import inspect_document
from document_tools.compare_ops import compare_documents
from document_tools.batch_ops import batch_process
from document_tools.ocr_ops import (
    _check_tesseract_or_raise,
    create_searchable_pdf,
    get_ocr_engine_status,
    ocr_pdf_document,
    ocr_single_image,
)
from document.ocr import (
    get_available_languages,
    get_default_ocr_lang,
    get_tesseract_config,
    resolve_ocr_lang,
    tesseract_available,
)
from fastapi import HTTPException

import pymupdf as fitz
from PIL import Image as PILImage
import docx
import pptx
import openpyxl


def _create_sample_pdf(pages=2, text="Sample document content"):
    doc = fitz.open()
    for i in range(pages):
        p = doc.new_page(width=595, height=842)
        p.insert_text((54, 100), f"Page {i+1}: {text}")
    buf = doc.tobytes()
    doc.close()
    return buf


def _create_sample_image(fmt="PNG", size=(200, 150)):
    im = PILImage.new("RGB", size, color=(34, 197, 94))
    buf = io.BytesIO()
    im.save(buf, format=fmt)
    return buf.getvalue()


def _create_sample_docx():
    doc = docx.Document()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("First paragraph with some text.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Value 1"
    table.rows[1].cells[1].text = "Value 2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _create_sample_pptx():
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Presentation Title"
    slide.placeholders[1].text = "Slide subtitle content"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _create_sample_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(["Item", "Quantity", "Price"])
    ws.append(["Server A", 5, 1200.50])
    ws.append(["Firewall B", 2, 850.00])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# 1. PDF TOOLS TESTS
# ═══════════════════════════════════════════════════════════════════════════

def test_reorder_pages():
    pdf = _create_sample_pdf(pages=3)
    reordered = reorder_pages(pdf, [3, 1, 2])
    doc = fitz.open(stream=reordered, filetype="pdf")
    assert len(doc) == 3
    assert "Page 3" in doc[0].get_text()
    assert "Page 1" in doc[1].get_text()
    doc.close()


def test_crop_pdf():
    pdf = _create_sample_pdf(pages=1)
    cropped = crop_pdf(pdf, margin_left=30, margin_top=40, margin_right=30, margin_bottom=40)
    doc = fitz.open(stream=cropped, filetype="pdf")
    rect = doc[0].rect
    assert rect.width < 595
    assert rect.height < 842
    doc.close()


def test_resize_pdf():
    pdf = _create_sample_pdf(pages=1)
    resized = resize_pdf(pdf, target_size="A3")
    doc = fitz.open(stream=resized, filetype="pdf")
    assert doc[0].rect.width == 842.0
    assert doc[0].rect.height == 1191.0
    doc.close()


def test_add_blank_page():
    pdf = _create_sample_pdf(pages=2)
    with_blank = add_blank_page(pdf, position="end")
    doc = fitz.open(stream=with_blank, filetype="pdf")
    assert len(doc) == 3
    doc.close()


def test_duplicate_page():
    pdf = _create_sample_pdf(pages=2)
    duplicated = duplicate_page(pdf, page_num=1, count=2)
    doc = fitz.open(stream=duplicated, filetype="pdf")
    assert len(doc) == 4
    doc.close()


def test_pdf_to_images_png_and_jpg():
    pdf = _create_sample_pdf(pages=2)
    png_imgs = render_pdf_to_images(pdf, fmt="PNG", pages="all")
    assert len(png_imgs) == 2
    assert "page_0001.png" in png_imgs
    assert png_imgs["page_0001.png"].startswith(b"\x89PNG")

    jpg_imgs = render_pdf_to_images(pdf, fmt="JPG", pages="1")
    assert len(jpg_imgs) == 1
    assert "page_0001.jpg" in jpg_imgs
    assert jpg_imgs["page_0001.jpg"][:3] == b"\xff\xd8\xff"


def test_compress_pdf_exact_metrics():
    pdf = _create_sample_pdf(pages=3)
    out_bytes, orig_size, out_size, pct = compress_pdf(pdf, level="medium")
    assert orig_size == len(pdf)
    assert out_size == len(out_bytes)
    assert isinstance(pct, float)
    assert not (pct is None)


# ═══════════════════════════════════════════════════════════════════════════
# 2. SECURITY TESTS
# ═══════════════════════════════════════════════════════════════════════════

def test_encrypt_and_decrypt_pdf():
    pdf = _create_sample_pdf(pages=1)
    enc = encrypt_pdf(pdf, user_password="secure_password_123")
    sec_info = inspect_pdf_security(enc)
    assert sec_info["encrypted"] is True

    dec = decrypt_pdf(enc, password="secure_password_123")
    dec_info = inspect_pdf_security(dec)
    assert dec_info["encrypted"] is False


def test_security_inspection_clean_pdf():
    pdf = _create_sample_pdf(pages=2)
    info = inspect_pdf_security(pdf)
    assert info["encrypted"] is False
    assert info["page_count"] == 2
    assert info["javascript_detected"] is False


# ═══════════════════════════════════════════════════════════════════════════
# 3. EDIT & SIGN TESTS
# ═══════════════════════════════════════════════════════════════════════════

def test_add_watermark():
    pdf = _create_sample_pdf(pages=2)
    wm = add_watermark(pdf, text="STRICTLY CONFIDENTIAL", opacity=0.4, angle=45)
    doc = fitz.open(stream=wm, filetype="pdf")
    assert "CONFIDENTIAL" in doc[0].get_text()
    doc.close()


def test_add_page_numbers():
    pdf = _create_sample_pdf(pages=2)
    numbered = add_page_numbers(pdf, position="bottom-center", template="Page {page} of {total}")
    doc = fitz.open(stream=numbered, filetype="pdf")
    assert "Page 1 of 2" in doc[0].get_text()
    assert "Page 2 of 2" in doc[1].get_text()
    doc.close()


def test_add_visual_signature():
    pdf = _create_sample_pdf(pages=1)
    sig_img = _create_sample_image("PNG", (100, 50))
    signed = add_visual_signature(pdf, sig_image_bytes=sig_img, page_num=1, x=200, y=500, width=150, height=60)
    doc = fitz.open(stream=signed, filetype="pdf")
    assert len(doc[0].get_images()) >= 1
    doc.close()


# ═══════════════════════════════════════════════════════════════════════════
# 4. OFFICE TESTS (DOCX, PPTX, XLSX)
# ═══════════════════════════════════════════════════════════════════════════

def test_docx_operations():
    docx_bytes = _create_sample_docx()
    text = docx_to_text(docx_bytes)
    assert "Test Heading" in text
    assert "First paragraph" in text

    md = docx_to_markdown(docx_bytes)
    assert "# Test Heading" in md

    tables = docx_extract_tables(docx_bytes)
    assert len(tables) == 1
    assert tables[0]["headers"] == ["Header A", "Header B"]

    stats = docx_get_stats(docx_bytes)
    assert stats["paragraph_count"] >= 2
    assert stats["table_count"] == 1

    pdf_out = docx_to_pdf(docx_bytes)
    assert pdf_out.startswith(b"%PDF")


def test_txt_conversions():
    raw_text = "Title: OFFAIR Document\n\nThis is paragraph one.\n\nThis is paragraph two."
    pdf_bytes = txt_to_pdf(raw_text, title="OFFAIR Title")
    assert pdf_bytes.startswith(b"%PDF")

    docx_bytes = txt_to_docx(raw_text, title="OFFAIR Title")
    assert docx_bytes.startswith(b"PK\x03\x04")


def test_pptx_operations():
    pptx_bytes = _create_sample_pptx()
    res = pptx_extract_text(pptx_bytes)
    assert res["slide_count"] == 1
    assert "Presentation Title" in res["slides"][0]["text"]

    stats = pptx_get_stats(pptx_bytes)
    assert stats["slide_count"] == 1
    assert stats["total_words"] > 0


def test_xlsx_operations():
    xlsx_bytes = _create_sample_xlsx()
    sheets_csv = xlsx_to_csv(xlsx_bytes)
    assert "Inventory.csv" in sheets_csv
    csv_str = sheets_csv["Inventory.csv"].decode("utf-8")
    assert "Server A" in csv_str

    preview = spreadsheet_preview(xlsx_bytes)
    assert preview["sheet_count"] == 1
    assert preview["sheets"][0]["name"] == "Inventory"

    xlsx_rebuilt = csv_to_xlsx(csv_str, sheet_name="Rebuilt")
    assert xlsx_rebuilt.startswith(b"PK\x03\x04")


# ═══════════════════════════════════════════════════════════════════════════
# 5. IMAGE OPERATIONS TESTS
# ═══════════════════════════════════════════════════════════════════════════

def test_image_transformations():
    img_png = _create_sample_image("PNG", (400, 300))
    resized, stats = resize_image(img_png, width=200, height=150)
    assert stats["output_dimensions"] == "200 × 150"

    converted, ext, cstats = convert_image(img_png, target_fmt="JPG")
    assert ext == "jpg"
    assert cstats["output_format"] == "JPG"

    compressed, cmp_stats = compress_image(img_png, quality=50)
    assert cmp_stats["output_size"] > 0

    rotated, rstats = rotate_flip_image(img_png, angle=90, flip_h=True)
    assert rstats["output_dimensions"] == "300 × 400"

    batch = batch_convert_images([(img_png, "pic1.png"), (img_png, "pic2.png")], target_fmt="WEBP")
    assert len(batch) == 2
    assert "pic1.webp" in batch


# ═══════════════════════════════════════════════════════════════════════════
# 6. DOCUMENT INSPECTOR & COMPARE TESTS
# ═══════════════════════════════════════════════════════════════════════════

def test_document_inspector():
    pdf_bytes = _create_sample_pdf(pages=2)
    res = inspect_document("test.pdf", pdf_bytes)
    assert res["category"] == "PDF Document"
    assert res["details"]["page_count"] == 2

    docx_bytes = _create_sample_docx()
    res_docx = inspect_document("sample.docx", docx_bytes)
    assert res_docx["category"] == "Word Document"

    xlsx_bytes = _create_sample_xlsx()
    res_xlsx = inspect_document("data.xlsx", xlsx_bytes)
    assert res_xlsx["category"] == "Excel Spreadsheet"


def test_compare_documents():
    t1 = "Line 1: Hello World\nLine 2: Sovereign AI\nLine 3: OFFAIR"
    t2 = "Line 1: Hello World\nLine 2: Modified Sovereign AI\nLine 3: OFFAIR"
    res = compare_documents(t1.encode("utf-8"), "doc1.txt", t2.encode("utf-8"), "doc2.txt")
    assert res["similarity_pct"] >= 60.0
    assert res["identical"] is False
    assert res["added_lines_count"] >= 1
    assert res["removed_lines_count"] >= 1


# ═══════════════════════════════════════════════════════════════════════════
# 7. BATCH PROCESSING TESTS (Isolation / Error Recovery)
# ═══════════════════════════════════════════════════════════════════════════

def test_batch_process_isolation():
    pdf1 = _create_sample_pdf(pages=1, text="PDF 1")
    pdf2 = _create_sample_pdf(pages=1, text="PDF 2")
    bad_file = b"This is not a PDF at all"

    # Batch process 3 files: 2 valid, 1 invalid.
    # Bad file must NOT crash the batch!
    files = [(pdf1, "doc1.pdf"), (bad_file, "bad.pdf"), (pdf2, "doc2.pdf")]
    outputs, summary = batch_process(files, operation="pdf-to-txt")

    assert summary["total_files"] == 3
    assert summary["completed"] == 2
    assert summary["failed"] == 1
    assert len(summary["errors"]) == 1
    assert summary["errors"][0]["file"] == "bad.pdf"
    assert "doc1.txt" in outputs
    assert "doc2.txt" in outputs


@pytest.mark.skipif(not tesseract_available(), reason="Tesseract missing")
def test_ocr_status():
    status = get_ocr_engine_status()
    assert "engine" in status
    assert "available" in status
    assert "cloud" in status
    assert status["cloud"] is False
    assert status["available"] is True
    assert "version" in status
    assert "languages" in status
    assert "eng" in status["languages"]


@pytest.mark.skipif(not tesseract_available(), reason="Tesseract missing")
def test_tesseract_config_and_language_resolution():
    cfg = get_tesseract_config()
    assert cfg["available"] is True
    assert "version" in cfg
    assert "command" in cfg
    langs = get_available_languages()
    assert "eng" in langs

    # Hindi should not be pretended to exist if hin is not in langs
    default_lang = get_default_ocr_lang()
    if "hin" in langs:
        assert default_lang == "eng+hin"
    else:
        assert default_lang == "eng"

    # resolve_ocr_lang falls back gracefully
    assert resolve_ocr_lang("nonexistent_lang") == default_lang
    assert resolve_ocr_lang("eng") == "eng"


@pytest.mark.skipif(not tesseract_available(), reason="Tesseract missing")
def test_ocr_single_image_with_real_file():
    img_path = Path(__file__).parent.parent / "demo_data" / "inspection_image.jpg"
    assert img_path.exists(), f"Image test file not found at {img_path}"

    with open(img_path, "rb") as f:
        img_bytes = f.read()

    extracted = ocr_single_image(img_bytes)
    assert len(extracted) > 20
    assert any(kw in extracted.upper() for kw in ["INSPECTION", "PUMP", "EQUIPMENT"])


@pytest.mark.skipif(not tesseract_available(), reason="Tesseract missing")
def test_ocr_pdf_document_with_real_file():
    pdf_path = Path(__file__).parent.parent / "demo_data" / "inspection_report.pdf"
    assert pdf_path.exists(), f"PDF test file not found at {pdf_path}"

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    res = ocr_pdf_document(pdf_bytes, pages="1")
    assert res["total_pages_ocr"] == 1
    assert len(res["pages"]) == 1
    page1_text = res["pages"][0]["text"]
    assert len(page1_text) > 20
    assert "INSPECTION" in page1_text.upper()
    assert res["full_text"]


@pytest.mark.skipif(not tesseract_available(), reason="Tesseract missing")
def test_create_searchable_pdf_with_real_file():
    pdf_path = Path(__file__).parent.parent / "demo_data" / "inspection_report.pdf"
    assert pdf_path.exists()

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    out_bytes = create_searchable_pdf(pdf_bytes, pages="1")
    assert len(out_bytes) > 0

    # Verify produced PDF is valid and contains text layer
    doc = fitz.open(stream=out_bytes, filetype="pdf")
    assert len(doc) >= 1
    text = doc[0].get_text()
    assert len(text) > 0
    doc.close()


def test_ocr_unavailable_error_handling(monkeypatch):
    import document_tools.ocr_ops as ocr_ops

    monkeypatch.setattr(ocr_ops, "tesseract_available", lambda: False)

    with pytest.raises(HTTPException) as exc_info:
        _check_tesseract_or_raise()

    assert exc_info.value.status_code == 503
    assert exc_info.value.detail == "Tesseract OCR is not available. Install Tesseract or configure TESSERACT_CMD."
