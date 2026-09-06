"""
document_tools/compare_ops.py
Deterministic offline document comparison using Python difflib.
Supports PDF vs PDF, DOCX vs DOCX, TXT vs TXT, and cross-format text diffing.
No AI/LLM required.
"""

from __future__ import annotations

import difflib
import io
from pathlib import Path
from typing import Any, Dict, List, Tuple

from fastapi import HTTPException

try:
    import pymupdf as fitz
    _PYMUPDF_AVAILABLE = True
except ImportError:
    try:
        import fitz
        _PYMUPDF_AVAILABLE = True
    except ImportError:
        _PYMUPDF_AVAILABLE = False

try:
    import docx
    from docx import Document
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

from document_tools.common import validate_file_size


def _extract_comparable_text(filename: str, data: bytes) -> List[str]:
    """Extract list of text lines from PDF, DOCX, or text file."""
    ext = Path(filename).suffix.lower().lstrip(".")

    # PDF
    if ext == "pdf" or data.startswith(b"%PDF"):
        if not _PYMUPDF_AVAILABLE:
            raise HTTPException(status_code=500, detail="PyMuPDF is required to compare PDF documents")
        doc = fitz.open(stream=data, filetype="pdf")
        lines = []
        for i, page in enumerate(doc):
            lines.append(f"--- Page {i + 1} ---")
            for line in page.get_text("text").splitlines():
                if line.strip():
                    lines.append(line.strip())
        doc.close()
        return lines

    # DOCX
    if ext in ("docx", "doc"):
        if not _DOCX_AVAILABLE or not data.startswith(b"PK\x03\x04"):
            raise HTTPException(status_code=400, detail="Cannot read DOCX document for comparison")
        doc = Document(io.BytesIO(data))
        lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())
        for t in doc.tables:
            for row in t.rows:
                row_str = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
                if row_str.strip():
                    lines.append(row_str)
        return lines

    # Plain text / fallback
    try:
        text = data.decode("utf-8", errors="replace")
        return [l.strip() for l in text.splitlines() if l.strip()]
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot extract text for comparison: {exc}")


def compare_documents(
    doc1_bytes: bytes,
    name1: str,
    doc2_bytes: bytes,
    name2: str,
) -> Dict[str, Any]:
    """
    Compare two documents line-by-line and return similarity ratio,
    added lines, removed lines, and unified diff output.
    """
    validate_file_size(doc1_bytes, label=name1)
    validate_file_size(doc2_bytes, label=name2)

    lines1 = _extract_comparable_text(name1, doc1_bytes)
    lines2 = _extract_comparable_text(name2, doc2_bytes)

    # Compute similarity ratio
    matcher = difflib.SequenceMatcher(None, lines1, lines2)
    similarity_pct = round(matcher.ratio() * 100.0, 1)

    # Unified diff
    diff_gen = difflib.unified_diff(
        lines1,
        lines2,
        fromfile=name1,
        tofile=name2,
        lineterm="",
    )
    diff_lines = list(diff_gen)

    added = []
    removed = []
    for line in diff_lines:
        if line.startswith("+") and not line.startswith("+++"):
            added.append(line[1:].strip())
        elif line.startswith("-") and not line.startswith("---"):
            removed.append(line[1:].strip())

    return {
        "file1": name1,
        "file2": name2,
        "similarity_pct": similarity_pct,
        "identical": similarity_pct == 100.0,
        "lines_count_file1": len(lines1),
        "lines_count_file2": len(lines2),
        "added_lines_count": len(added),
        "removed_lines_count": len(removed),
        "sample_added": added[:20],
        "sample_removed": removed[:20],
        "diff_text": "\n".join(diff_lines[:200]),
    }
