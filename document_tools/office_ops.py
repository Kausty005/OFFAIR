"""
document_tools/office_ops.py
Purely local Office document processing (Word, PowerPoint, Excel).
Uses python-docx, python-pptx, openpyxl, and PyMuPDF.
Air-gap compliant, zero cloud dependencies.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException

try:
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

try:
    import pptx
    _PPTX_AVAILABLE = True
except ImportError:
    _PPTX_AVAILABLE = False

try:
    import openpyxl
    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False

try:
    import pymupdf as fitz
    _PYMUPDF_AVAILABLE = True
except ImportError:
    try:
        import fitz
        _PYMUPDF_AVAILABLE = True
    except ImportError:
        _PYMUPDF_AVAILABLE = False

from document_tools.common import (
    validate_docx,
    validate_pptx,
    validate_xlsx,
)


# ─────────────────────────────────────────────────────────────────────────────
# 1. WORD / DOCX OPERATIONS
# ─────────────────────────────────────────────────────────────────────────────

def docx_to_text(data: bytes) -> str:
    """Extract plain text from a Word (.docx) document."""
    validate_docx(data)
    if not _DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx is not installed on this server")

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot read DOCX document: {exc}")

    lines: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)

    # Also extract text from tables
    if doc.tables:
        lines.append("\n--- Table Data ---")
        for t_idx, table in enumerate(doc.tables):
            lines.append(f"[Table {t_idx + 1}]")
            for row in table.rows:
                row_vals = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                lines.append(" | ".join(row_vals))
            lines.append("")

    return "\n\n".join(lines)


def docx_to_markdown(data: bytes) -> str:
    """Extract structured Markdown with headings, bullet lists, and tables."""
    validate_docx(data)
    if not _DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx is not installed on this server")

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot read DOCX document: {exc}")

    md_lines: List[str] = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style_name = p.style.name.lower() if p.style else ""
        if "heading 1" in style_name:
            md_lines.append(f"# {txt}\n")
        elif "heading 2" in style_name:
            md_lines.append(f"## {txt}\n")
        elif "heading 3" in style_name:
            md_lines.append(f"### {txt}\n")
        elif "list" in style_name or "bullet" in style_name:
            md_lines.append(f"- {txt}")
        else:
            md_lines.append(f"{txt}\n")

    for t_idx, table in enumerate(doc.tables):
        md_lines.append(f"\n### Table {t_idx + 1}\n")
        rows = table.rows
        if not rows:
            continue
        headers = [c.text.strip().replace("\n", " ") for c in rows[0].cells]
        md_lines.append("| " + " | ".join(headers) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in rows[1:]:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            md_lines.append("| " + " | ".join(cells) + " |")
        md_lines.append("")

    return "\n".join(md_lines)


def docx_extract_tables(data: bytes) -> List[Dict[str, Any]]:
    """Extract all tables from a Word document as JSON structures."""
    validate_docx(data)
    if not _DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx is not installed on this server")

    doc = Document(io.BytesIO(data))
    tables_data: List[Dict[str, Any]] = []

    for idx, table in enumerate(doc.tables):
        rows_data: List[List[str]] = []
        for row in table.rows:
            rows_data.append([c.text.strip() for c in row.cells])
        headers = rows_data[0] if rows_data else []
        content_rows = rows_data[1:] if len(rows_data) > 1 else []
        tables_data.append({
            "table_index": idx + 1,
            "headers": headers,
            "rows": content_rows,
            "row_count": len(rows_data),
            "col_count": len(headers),
        })

    return tables_data


def docx_get_stats(data: bytes) -> Dict[str, Any]:
    """Calculate word, character, paragraph, and table metrics from DOCX."""
    validate_docx(data)
    if not _DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx is not installed on this server")

    doc = Document(io.BytesIO(data))
    text_content = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    words = re.findall(r"\b\w+\b", text_content)

    meta = doc.core_properties
    return {
        "paragraph_count": len([p for p in doc.paragraphs if p.text.strip()]),
        "word_count": len(words),
        "char_count": len(text_content),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "title": meta.title or "—",
        "author": meta.author or "—",
        "created": str(meta.created) if meta.created else "—",
        "modified": str(meta.modified) if meta.modified else "—",
    }


def docx_to_pdf(data: bytes) -> bytes:
    """
    Render DOCX document into a clean, paginated A4 PDF.
    Pure local conversion using python-docx and PyMuPDF layout.
    """
    validate_docx(data)
    if not _DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx is not installed on this server")
    if not _PYMUPDF_AVAILABLE:
        raise HTTPException(status_code=500, detail="PyMuPDF is not installed on this server")

    doc = Document(io.BytesIO(data))
    pdf_doc = fitz.open()

    page_w, page_h = 595.0, 842.0
    margin_x = 54.0
    margin_y = 54.0
    usable_w = page_w - (2 * margin_x)
    max_y = page_h - margin_y

    current_page = pdf_doc.new_page(width=page_w, height=page_h)
    cursor_y = margin_y

    def ensure_space(height_needed: float) -> None:
        nonlocal current_page, cursor_y
        if cursor_y + height_needed > max_y:
            current_page = pdf_doc.new_page(width=page_w, height=page_h)
            cursor_y = margin_y

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            cursor_y += 8
            continue

        style_name = p.style.name.lower() if p.style else ""
        if "heading 1" in style_name:
            font_size = 18.0
            color = (0.1, 0.1, 0.2)
            space_after = 12.0
        elif "heading 2" in style_name:
            font_size = 14.0
            color = (0.2, 0.2, 0.3)
            space_after = 10.0
        elif "heading 3" in style_name:
            font_size = 12.0
            color = (0.3, 0.3, 0.3)
            space_after = 8.0
        else:
            font_size = 10.0
            color = (0.1, 0.1, 0.1)
            space_after = 6.0

        # Estimate text height and word wrap
        line_height = font_size * 1.35
        # Approximate chars per line
        chars_per_line = int(usable_w / (font_size * 0.5))
        words = txt.split()
        wrapped_lines: List[str] = []
        cur_line = []
        cur_len = 0
        for w in words:
            if cur_len + len(w) + 1 > chars_per_line:
                wrapped_lines.append(" ".join(cur_line))
                cur_line = [w]
                cur_len = len(w)
            else:
                cur_line.append(w)
                cur_len += len(w) + 1
        if cur_line:
            wrapped_lines.append(" ".join(cur_line))

        block_height = len(wrapped_lines) * line_height + space_after
        ensure_space(block_height)

        for line_str in wrapped_lines:
            current_page.insert_text(
                (margin_x, cursor_y + font_size),
                line_str,
                fontsize=font_size,
                color=color,
            )
            cursor_y += line_height
        cursor_y += space_after

    # Render tables
    for table in doc.tables:
        col_count = len(table.columns)
        if col_count == 0:
            continue
        col_w = usable_w / col_count
        row_h = 22.0

        ensure_space(row_h * (len(table.rows) + 1))

        for r_idx, row in enumerate(table.rows):
            ensure_space(row_h)
            is_header = (r_idx == 0)
            for c_idx, cell in enumerate(row.cells):
                rect = fitz.Rect(
                    margin_x + (c_idx * col_w),
                    cursor_y,
                    margin_x + ((c_idx + 1) * col_w),
                    cursor_y + row_h,
                )
                current_page.draw_rect(rect, color=(0.7, 0.7, 0.7), width=0.5)
                cell_text = cell.text.strip().replace("\n", " ")[:30]
                current_page.insert_text(
                    (rect.x0 + 4, rect.y0 + 15),
                    cell_text,
                    fontsize=9.0,
                    color=(0.0, 0.0, 0.0) if is_header else (0.2, 0.2, 0.2),
                )
            cursor_y += row_h
        cursor_y += 14

    out_bytes = pdf_doc.tobytes(garbage=4, deflate=True)
    pdf_doc.close()
    return out_bytes


def txt_to_docx(text: str, title: str = "Document") -> bytes:
    """Convert raw text into a styled Word (.docx) document."""
    if not _DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx is not installed on this server")

    doc = Document()
    doc.add_heading(title, level=1)

    for para in text.split("\n\n"):
        clean = para.strip()
        if clean:
            doc.add_paragraph(clean)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def txt_to_pdf(text: str, title: str = "Document") -> bytes:
    """Convert raw text into a clean paginated A4 PDF."""
    if not _PYMUPDF_AVAILABLE:
        raise HTTPException(status_code=500, detail="PyMuPDF is not installed on this server")

    pdf_doc = fitz.open()
    page_w, page_h = 595.0, 842.0
    margin_x, margin_y = 54.0, 54.0
    usable_w = page_w - (2 * margin_x)
    max_y = page_h - margin_y

    current_page = pdf_doc.new_page(width=page_w, height=page_h)
    cursor_y = margin_y

    # Title
    current_page.insert_text((margin_x, cursor_y + 18), title, fontsize=18, color=(0.1, 0.1, 0.2))
    cursor_y += 36

    font_size = 10.0
    line_height = 14.0
    chars_per_line = int(usable_w / (font_size * 0.5))

    for line in text.splitlines():
        if not line.strip():
            cursor_y += 8
            continue

        words = line.split()
        wrapped_lines: List[str] = []
        cur_line = []
        cur_len = 0
        for w in words:
            if cur_len + len(w) + 1 > chars_per_line:
                wrapped_lines.append(" ".join(cur_line))
                cur_line = [w]
                cur_len = len(w)
            else:
                cur_line.append(w)
                cur_len += len(w) + 1
        if cur_line:
            wrapped_lines.append(" ".join(cur_line))

        for w_line in wrapped_lines:
            if cursor_y + line_height > max_y:
                current_page = pdf_doc.new_page(width=page_w, height=page_h)
                cursor_y = margin_y
            current_page.insert_text((margin_x, cursor_y + font_size), w_line, fontsize=font_size, color=(0.1, 0.1, 0.1))
            cursor_y += line_height

    out_bytes = pdf_doc.tobytes(garbage=4, deflate=True)
    pdf_doc.close()
    return out_bytes


# ─────────────────────────────────────────────────────────────────────────────
# 2. PRESENTATION / PPTX OPERATIONS
# ─────────────────────────────────────────────────────────────────────────────

def pptx_extract_text(data: bytes) -> Dict[str, Any]:
    """Extract slide text, slide notes, and stats from PPTX."""
    validate_pptx(data)
    if not _PPTX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-pptx is not installed on this server")

    try:
        prs = pptx.Presentation(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot read PPTX presentation: {exc}")

    slides_info: List[Dict[str, Any]] = []
    total_words = 0

    for i, slide in enumerate(prs.slides):
        slide_text_parts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        slide_text_parts.append(t)

        notes_text = ""
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf:
                notes_text = notes_tf.text.strip()

        combined_text = "\n".join(slide_text_parts)
        words_in_slide = len(re.findall(r"\b\w+\b", combined_text))
        total_words += words_in_slide

        slides_info.append({
            "slide_number": i + 1,
            "text": combined_text,
            "notes": notes_text,
            "word_count": words_in_slide,
            "shape_count": len(slide.shapes),
        })

    return {
        "slide_count": len(prs.slides),
        "total_words": total_words,
        "slides": slides_info,
    }


def pptx_get_stats(data: bytes) -> Dict[str, Any]:
    """Quick presentation statistics."""
    info = pptx_extract_text(data)
    has_notes_count = sum(1 for s in info["slides"] if s["notes"])
    return {
        "slide_count": info["slide_count"],
        "total_words": info["total_words"],
        "slides_with_notes": has_notes_count,
    }


# ─────────────────────────────────────────────────────────────────────────────
# 3. SPREADSHEET / XLSX & CSV OPERATIONS
# ─────────────────────────────────────────────────────────────────────────────

def xlsx_to_csv(data: bytes, sheet_name: Optional[str] = None) -> Dict[str, bytes]:
    """Convert Excel worksheet(s) to CSV byte streams."""
    validate_xlsx(data)
    if not _OPENPYXL_AVAILABLE:
        raise HTTPException(status_code=500, detail="openpyxl is not installed on this server")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot read XLSX workbook: {exc}")

    results: Dict[str, bytes] = {}
    target_sheets = [sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.sheetnames

    for s_name in target_sheets:
        ws = wb[s_name]
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            cleaned_row = ["" if cell is None else str(cell) for cell in row]
            # Ignore purely empty rows
            if any(cleaned_row):
                writer.writerow(cleaned_row)
        clean_filename = f"{s_name}.csv".replace(" ", "_")
        results[clean_filename] = buf.getvalue().encode("utf-8")

    return results


def csv_to_xlsx(csv_text: str, sheet_name: str = "Data") -> bytes:
    """Convert CSV string into a styled XLSX workbook."""
    if not _OPENPYXL_AVAILABLE:
        raise HTTPException(status_code=500, detail="openpyxl is not installed on this server")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name or "Sheet1"

    reader = csv.reader(io.StringIO(csv_text))
    for r_idx, row in enumerate(reader, start=1):
        ws.append(row)
        if r_idx == 1:
            # Style header row with bold text
            for cell in ws[1]:
                cell.font = openpyxl.styles.Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def spreadsheet_preview(data: bytes) -> Dict[str, Any]:
    """Extract sheet names, dimensions, and first 10 rows for preview."""
    validate_xlsx(data)
    if not _OPENPYXL_AVAILABLE:
        raise HTTPException(status_code=500, detail="openpyxl is not installed on this server")

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets_info = []

    for name in wb.sheetnames:
        ws = wb[name]
        preview_rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 10:
                break
            preview_rows.append([str(c) if c is not None else "" for c in row])

        sheets_info.append({
            "name": name,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "preview": preview_rows,
        })

    return {
        "sheet_count": len(wb.sheetnames),
        "sheets": sheets_info,
    }
