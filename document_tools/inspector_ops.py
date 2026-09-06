"""
document_tools/inspector_ops.py
Deterministic local document inspector.
Analyzes PDFs, Word documents, PowerPoint presentations, Excel spreadsheets, images, and text files.
No external network calls, completely offline.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import HTTPException
from PIL import Image as PILImage

try:
    import pymupdf as fitz
    _PYMUPDF_AVAILABLE = True
except ImportError:
    try:
        import fitz
        _PYMUPDF_AVAILABLE = True
    except ImportError:
        _PYMUPDF_AVAILABLE = False

try:
    import docx
    from docx import Document
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

try:
    import pptx
    _PPTX_AVAILABLE = True
except ImportError:
    _PPTX_AVAILABLE = False

try:
    import openpyxl
    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False

from document_tools.common import safe_filename, validate_file_size


def _format_size(num_bytes: int) -> str:
    if num_bytes < 1024:
        return f"{num_bytes} B"
    elif num_bytes < 1024 * 1024:
        return f"{num_bytes / 1024:.1f} KB"
    else:
        return f"{num_bytes / (1024 * 1024):.2f} MB"


def inspect_document(filename: str, data: bytes) -> Dict[str, Any]:
    """Inspect any supported document and return detailed structural metrics."""
    validate_file_size(data, label=filename)
    safe_name = safe_filename(filename)
    ext = Path(filename).suffix.lower().lstrip(".")
    size_bytes = len(data)

    base_result: Dict[str, Any] = {
        "filename": safe_name,
        "file_size": size_bytes,
        "formatted_size": _format_size(size_bytes),
        "extension": ext.upper() if ext else "UNKNOWN",
        "category": "Unknown",
        "details": {},
    }

    # 1. PDF
    if ext == "pdf" or data.startswith(b"%PDF"):
        base_result["category"] = "PDF Document"
        if not _PYMUPDF_AVAILABLE:
            base_result["details"]["note"] = "PyMuPDF not available for inspection"
            return base_result

        doc = fitz.open(stream=data, filetype="pdf")
        total_pages = len(doc)
        is_encrypted = doc.is_encrypted
        image_count = 0
        total_chars = 0
        total_words = 0

        if not is_encrypted:
            for page in doc:
                text = page.get_text("text")
                total_chars += len(text)
                words = re.findall(r"\b\w+\b", text)
                total_words += len(words)
                image_count += len(page.get_images())

        meta = doc.metadata or {}
        doc.close()

        base_result["details"] = {
            "page_count": total_pages,
            "encrypted": is_encrypted,
            "word_count": total_words if not is_encrypted else "Protected",
            "char_count": total_chars if not is_encrypted else "Protected",
            "image_count": image_count if not is_encrypted else "Protected",
            "title": meta.get("title") or "—",
            "author": meta.get("author") or "—",
            "producer": meta.get("producer") or "—",
            "creation_date": meta.get("creationDate") or "—",
            "modification_date": meta.get("modDate") or "—",
        }
        return base_result

    # 2. DOCX
    if ext in ("docx", "doc"):
        base_result["category"] = "Word Document"
        if not _DOCX_AVAILABLE or not data.startswith(b"PK\x03\x04"):
            base_result["details"]["note"] = "Legacy DOC or missing python-docx"
            return base_result

        try:
            w_doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in w_doc.paragraphs if p.text.strip())
            words = re.findall(r"\b\w+\b", text)
            core = w_doc.core_properties
            base_result["details"] = {
                "paragraph_count": len([p for p in w_doc.paragraphs if p.text.strip()]),
                "word_count": len(words),
                "char_count": len(text),
                "table_count": len(w_doc.tables),
                "section_count": len(w_doc.sections),
                "title": core.title or "—",
                "author": core.author or "—",
                "created": str(core.created) if core.created else "—",
                "modified": str(core.modified) if core.modified else "—",
            }
        except Exception as exc:
            base_result["details"]["error"] = f"Failed to parse DOCX: {exc}"
        return base_result

    # 3. PPTX
    if ext in ("pptx", "ppt"):
        base_result["category"] = "PowerPoint Presentation"
        if not _PPTX_AVAILABLE or not data.startswith(b"PK\x03\x04"):
            base_result["details"]["note"] = "Legacy PPT or missing python-pptx"
            return base_result

        try:
            prs = pptx.Presentation(io.BytesIO(data))
            slide_count = len(prs.slides)
            total_words = 0
            shapes_count = 0
            notes_count = 0
            for s in prs.slides:
                shapes_count += len(s.shapes)
                if s.has_notes_slide and s.notes_slide.notes_text_frame:
                    if s.notes_slide.notes_text_frame.text.strip():
                        notes_count += 1
                for shp in s.shapes:
                    if shp.has_text_frame:
                        for p in shp.text_frame.paragraphs:
                            words = re.findall(r"\b\w+\b", p.text)
                            total_words += len(words)

            base_result["details"] = {
                "slide_count": slide_count,
                "total_words": total_words,
                "shape_count": shapes_count,
                "slides_with_notes": notes_count,
            }
        except Exception as exc:
            base_result["details"]["error"] = f"Failed to parse PPTX: {exc}"
        return base_result

    # 4. XLSX
    if ext in ("xlsx", "xls"):
        base_result["category"] = "Excel Spreadsheet"
        if not _OPENPYXL_AVAILABLE or not data.startswith(b"PK\x03\x04"):
            base_result["details"]["note"] = "Legacy XLS or missing openpyxl"
            return base_result

        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            total_rows = 0
            total_cols = 0
            for name in wb.sheetnames:
                ws = wb[name]
                total_rows += (ws.max_row or 0)
                total_cols = max(total_cols, ws.max_column or 0)

            base_result["details"] = {
                "sheet_count": len(wb.sheetnames),
                "sheet_names": wb.sheetnames,
                "total_rows": total_rows,
                "max_columns": total_cols,
            }
        except Exception as exc:
            base_result["details"]["error"] = f"Failed to parse XLSX: {exc}"
        return base_result

    # 5. Image
    if ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp") or data.startswith(b"\x89PNG") or data[:3] == b"\xff\xd8\xff":
        base_result["category"] = "Raster Image"
        try:
            img = PILImage.open(io.BytesIO(data))
            base_result["details"] = {
                "format": img.format or ext.upper(),
                "dimensions": f"{img.width} × {img.height}",
                "width": img.width,
                "height": img.height,
                "color_mode": img.mode,
                "animated": getattr(img, "is_animated", False),
            }
        except Exception as exc:
            base_result["details"]["error"] = f"Failed to parse image: {exc}"
        return base_result

    # 6. Plain Text / Markdown / CSV
    try:
        txt = data.decode("utf-8")
        words = re.findall(r"\b\w+\b", txt)
        lines = txt.splitlines()
        base_result["category"] = "Text Document"
        base_result["details"] = {
            "line_count": len(lines),
            "word_count": len(words),
            "char_count": len(txt),
            "encoding": "UTF-8",
        }
        return base_result
    except UnicodeDecodeError:
        base_result["category"] = "Binary File"
        base_result["details"] = {
            "raw_size": size_bytes,
            "note": "Non-text binary file",
        }
        return base_result
