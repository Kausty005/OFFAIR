"""
tools/docx_generator.py
Professional Word document generator using python-docx.
Creates structured industrial documents with headers, tables, and footers.
"""

import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

_base = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, _base)

from security.audit import log

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False


def _add_heading(doc, text: str, level: int = 1):
    heading = doc.add_heading(text, level=level)
    return heading


def _add_table(doc, headers: list[str], rows: list[list[str]]):
    """Add a styled table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Shading Accent 1"

    # Header row
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.bold = True

    # Data rows
    for row_idx, row in enumerate(rows):
        cells = table.rows[row_idx + 1].cells
        for col_idx, val in enumerate(row):
            cells[col_idx].text = str(val)

    return table


def _add_footer(doc, text: str):
    """Add footer to all sections."""
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def create_maintenance_approval_note(
    equipment_name: str,
    equipment_id: str,
    inspection_date: str,
    findings: list[str],
    measurements: dict,
    recommendations: str,
    sop_references: list[dict],
    ai_reasoning: str,
    output_path: str | Path,
    risk_level: str = "HIGH",
) -> dict:
    """
    Generate a professional Maintenance Approval Note DOCX.

    Returns:
        {"success": bool, "path": str, "error": str|None}
    """
    if not _DOCX_AVAILABLE:
        return {"success": False, "path": "", "error": "python-docx not installed. Run: pip install python-docx"}

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = Document()

        # ── Document Metadata ──────────────────────────────────────────────
        doc.core_properties.author = "Sovereign AI Workbench — SIH26117"
        doc.core_properties.title = "Maintenance Approval Note"
        doc.core_properties.subject = f"Equipment: {equipment_id}"

        # ── Title ──────────────────────────────────────────────────────────
        title = doc.add_heading("MAINTENANCE APPROVAL NOTE", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run("⚠️  AI-GENERATED DRAFT — REQUIRES AUTHORIZED HUMAN REVIEW")
        run.bold = True

        doc.add_paragraph()  # Spacer

        # ── 1. Subject ─────────────────────────────────────────────────────
        _add_heading(doc, "1. Subject", 2)
        doc.add_paragraph(
            f"Request for maintenance approval for equipment {equipment_id} — {equipment_name} "
            f"based on inspection findings dated {inspection_date}."
        )

        # ── 2. Equipment Details ───────────────────────────────────────────
        _add_heading(doc, "2. Equipment Details", 2)
        _add_table(
            doc,
            headers=["Parameter", "Value"],
            rows=[
                ["Equipment Name", equipment_name],
                ["Equipment ID", equipment_id],
                ["Inspection Date", inspection_date],
                ["Risk Level", risk_level],
                ["Document Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ],
        )

        # ── 3. Inspection Summary ──────────────────────────────────────────
        _add_heading(doc, "3. Inspection Summary", 2)
        doc.add_paragraph(
            f"A detailed inspection of {equipment_name} ({equipment_id}) was conducted on {inspection_date}. "
            "The following observations were recorded and processed through the Sovereign AI Workbench (SIH26117) "
            "for automated analysis and approval note generation."
        )

        # ── 4. Key Findings ────────────────────────────────────────────────
        _add_heading(doc, "4. Key Findings", 2)
        if findings:
            for finding in findings:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(finding)
        else:
            doc.add_paragraph("No specific findings extracted.")

        # ── 5. Measurements ───────────────────────────────────────────────
        _add_heading(doc, "5. Measurements & Observations", 2)
        if measurements:
            meas_rows = [[k, str(v)] for k, v in measurements.items()]
            _add_table(doc, headers=["Parameter", "Measured Value"], rows=meas_rows)
        else:
            doc.add_paragraph("No specific measurements recorded.")

        # ── 6. SOP/Guideline References ────────────────────────────────────
        _add_heading(doc, "6. Relevant SOP / Guideline References", 2)
        if sop_references:
            for ref in sop_references:
                doc_name = ref.get("document", "")
                page = ref.get("page", "")
                snippet = ref.get("text", "")[:200]
                p = doc.add_paragraph(style="List Bullet")
                label = f"{doc_name}" + (f" — Page {page}" if page else "")
                p.add_run(label).bold = True
                if snippet:
                    doc.add_paragraph(f'   "{snippet}..."')
        else:
            doc.add_paragraph("No SOP references retrieved. Manual SOP lookup recommended.")

        # ── 7. Risk / Operational Considerations ──────────────────────────
        _add_heading(doc, "7. Risk / Operational Considerations", 2)
        risk_text = {
            "CRITICAL": "Immediate shutdown and maintenance required. Risk of equipment failure and safety incident.",
            "HIGH": "Maintenance should be scheduled at earliest opportunity. Monitor equipment closely in the interim.",
            "MEDIUM": "Schedule maintenance in next planned shutdown window. Continue monitoring.",
            "LOW": "Minor maintenance recommended during next scheduled service.",
        }.get(risk_level.upper(), "Consult maintenance engineer for risk assessment.")
        doc.add_paragraph(risk_text)

        # ── 8. Recommended Action ─────────────────────────────────────────
        _add_heading(doc, "8. Recommended Action", 2)
        doc.add_paragraph(recommendations or "Detailed inspection and maintenance by qualified engineer.")

        # ── 9. Approval Requested ─────────────────────────────────────────
        _add_heading(doc, "9. Approval Requested", 2)
        doc.add_paragraph(
            "Authorization is requested to proceed with the recommended maintenance action on "
            f"{equipment_name} ({equipment_id})."
        )
        doc.add_paragraph()
        _add_table(
            doc,
            headers=["Role", "Name", "Signature", "Date"],
            rows=[
                ["Maintenance Engineer", "", "_______________", ""],
                ["Section Head", "", "_______________", ""],
                ["Plant Manager", "", "_______________", ""],
            ],
        )

        # ── 10. AI Processing Audit Information ───────────────────────────
        _add_heading(doc, "10. AI Processing Audit Information", 2)
        _add_table(
            doc,
            headers=["Parameter", "Value"],
            rows=[
                ["AI System", "Sovereign AI Workbench (SIH26117)"],
                ["Inference", "LOCAL — No cloud API"],
                ["OCR", "LOCAL — Tesseract / Vision Model"],
                ["RAG", f"LOCAL — {len(sop_references)} sources retrieved"],
                ["External API Calls", "0"],
                ["Generated At", datetime.now().isoformat()],
                ["Human Review Required", "YES — This is a draft"],
            ],
        )

        doc.add_paragraph()
        disclaimer = doc.add_paragraph()
        run = disclaimer.add_run(
            "DISCLAIMER: This document was generated by an AI system as a draft for human review. "
            "All findings and recommendations must be verified by qualified personnel before any maintenance action is taken. "
            "The AI system does not have authority to approve maintenance operations."
        )
        run.italic = True

        # ── Footer ────────────────────────────────────────────────────────
        _add_footer(
            doc,
            f"SIH26117 Sovereign AI Workbench | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | LOCAL | DRAFT — HUMAN REVIEW REQUIRED"
        )

        # Save
        doc.save(str(output_path))
        log("DOCX_GENERATED", path=str(output_path), size=output_path.stat().st_size)

        return {
            "success": True,
            "path": str(output_path),
            "size": output_path.stat().st_size,
            "error": None,
        }

    except Exception as e:
        log("DOCX_ERROR", error=str(e))
        return {"success": False, "path": "", "error": str(e)}


def create_coding_report(
    prompt: str,
    code: str,
    test_code: str,
    sandbox_result: dict,
    output_path: str | Path,
) -> dict:
    """Generate a coding agent report DOCX."""
    if not _DOCX_AVAILABLE:
        return {"success": False, "path": "", "error": "python-docx not installed"}

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = Document()
        doc.add_heading("CODING AGENT REPORT", 0)
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_heading("1. Task", 2)
        doc.add_paragraph(prompt)
        doc.add_heading("2. Generated Code", 2)
        doc.add_paragraph(code, style="No Spacing")
        doc.add_heading("3. Test Cases", 2)
        doc.add_paragraph(test_code or "No test cases.", style="No Spacing")
        doc.add_heading("4. Sandbox Execution Results", 2)
        _add_table(
            doc,
            headers=["Parameter", "Value"],
            rows=[
                ["Sandbox", "Docker"],
                ["Network", "DISABLED"],
                ["Tests Passed", str(sandbox_result.get("tests_passed", 0))],
                ["Tests Total", str(sandbox_result.get("tests_total", 0))],
                ["Success", str(sandbox_result.get("success", False))],
                ["Elapsed", f"{sandbox_result.get('elapsed_s', 0)}s"],
            ],
        )
        _add_footer(doc, f"SIH26117 Coding Agent | {datetime.now().strftime('%Y-%m-%d %H:%M')} | LOCAL EXECUTION")
        doc.save(str(output_path))
        log("DOCX_GENERATED", path=str(output_path), type="coding")
        return {"success": True, "path": str(output_path), "error": None}
    except Exception as e:
        return {"success": False, "path": "", "error": str(e)}


def docx_available() -> bool:
    return _DOCX_AVAILABLE
